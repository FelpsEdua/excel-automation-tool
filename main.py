import pandas as pd
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


# ==========================================================
# CARREGAR ARQUIVO
# ==========================================================

def carregar_arquivo(caminho):

    arquivo = Path(caminho)

    if not arquivo.exists():
        raise FileNotFoundError(
            f"Arquivo não encontrado: {caminho}"
        )

    extensao = arquivo.suffix.lower()

    if extensao == ".csv":
        return pd.read_csv(arquivo)

    if extensao in [".xlsx", ".xls"]:
        return pd.read_excel(arquivo)

    raise ValueError(
        "Formato não suportado. "
        "Use CSV ou Excel."
    )


# ==========================================================
# LIMPAR DADOS
# ==========================================================

def limpar_dados(df):

    df = df.copy()

    registros_originais = len(df)

    # Remove linhas completamente vazias
    df = df.dropna(how="all")

    # Limpa nomes das colunas
    df.columns = (
        df.columns
        .astype(str)
        .str.strip()
    )

    # Limpa espaços em textos
    for coluna in df.select_dtypes(
        include=["object", "string"]
    ).columns:

        df[coluna] = df[coluna].apply(
            lambda valor:
            valor.strip()
            if isinstance(valor, str)
            else valor
        )

    # Conta duplicados
    duplicados = df.duplicated().sum()

    # Remove duplicados
    df = df.drop_duplicates()

    # Preenche vazios
    for coluna in df.columns:

        if df[coluna].dtype == "object":

            df[coluna] = df[coluna].fillna(
                "Não informado"
            )

        else:

            df[coluna] = df[coluna].fillna(0)

    return (
        df,
        registros_originais,
        duplicados
    )


# ==========================================================
# FATURAMENTO
# ==========================================================

def calcular_faturamento(df):

    if {
        "Quantidade",
        "Valor"
    }.issubset(df.columns):

        quantidade = pd.to_numeric(
            df["Quantidade"],
            errors="coerce"
        ).fillna(0)

        valor = pd.to_numeric(
            df["Valor"],
            errors="coerce"
        ).fillna(0)

        return (
            quantidade * valor
        ).sum()

    return 0


# ==========================================================
# CRIAR RELATÓRIO
# ==========================================================

def criar_relatorio(
    df,
    registros_originais,
    duplicados,
    caminho_saida
):

    caminho_saida = Path(caminho_saida)

    caminho_saida.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    faturamento_total = calcular_faturamento(df)

    with pd.ExcelWriter(
        caminho_saida,
        engine="openpyxl"
    ) as writer:

        # ----------------------------------------------
        # DADOS LIMPOS
        # ----------------------------------------------

        df.to_excel(
            writer,
            sheet_name="Dados Limpos",
            index=False
        )

        # ----------------------------------------------
        # RESUMO TÉCNICO
        # ----------------------------------------------

        resumo = pd.DataFrame({

            "Indicador": [
                "Registros originais",
                "Registros finais",
                "Registros duplicados removidos",
                "Total de colunas",
                "Faturamento total"
            ],

            "Valor": [
                registros_originais,
                len(df),
                duplicados,
                len(df.columns),
                faturamento_total
            ]
        })

        resumo.to_excel(
            writer,
            sheet_name="Resumo",
            index=False
        )

        # ----------------------------------------------
        # POR PRODUTO
        # ----------------------------------------------

        if {
            "Produto",
            "Quantidade",
            "Valor"
        }.issubset(df.columns):

            dados = df.copy()

            dados["Quantidade"] = pd.to_numeric(
                dados["Quantidade"],
                errors="coerce"
            ).fillna(0)

            dados["Valor"] = pd.to_numeric(
                dados["Valor"],
                errors="coerce"
            ).fillna(0)

            dados["Faturamento"] = (
                dados["Quantidade"]
                * dados["Valor"]
            )

            por_produto = (
                dados
                .groupby(
                    "Produto",
                    as_index=False
                )
                .agg({
                    "Quantidade": "sum",
                    "Faturamento": "sum"
                })
                .sort_values(
                    "Faturamento",
                    ascending=False
                )
            )

            por_produto.to_excel(
                writer,
                sheet_name="Por Produto",
                index=False
            )

        # ----------------------------------------------
        # POR CLIENTE
        # ----------------------------------------------

        if {
            "Cliente",
            "Quantidade",
            "Valor"
        }.issubset(df.columns):

            dados = df.copy()

            dados["Quantidade"] = pd.to_numeric(
                dados["Quantidade"],
                errors="coerce"
            ).fillna(0)

            dados["Valor"] = pd.to_numeric(
                dados["Valor"],
                errors="coerce"
            ).fillna(0)

            dados["Faturamento"] = (
                dados["Quantidade"]
                * dados["Valor"]
            )

            por_cliente = (
                dados
                .groupby(
                    "Cliente",
                    as_index=False
                )
                .agg({
                    "Quantidade": "sum",
                    "Faturamento": "sum"
                })
                .sort_values(
                    "Faturamento",
                    ascending=False
                )
            )

            por_cliente.to_excel(
                writer,
                sheet_name="Por Cliente",
                index=False
            )

    return caminho_saida


# ==========================================================
# FORMATAÇÃO DAS ABAS
# ==========================================================

def formatar_abas(wb):

    azul = "1F4E78"
    branco = "FFFFFF"

    borda = Border(
        bottom=Side(
            style="thin",
            color="D9E1F2"
        )
    )

    for ws in wb.worksheets:

        # Não formatamos o dashboard aqui
        if ws.title == "Resumo":
            continue

        # ----------------------------------------------
        # CABEÇALHO
        # ----------------------------------------------

        for cell in ws[1]:

            cell.font = Font(
                bold=True,
                color=branco
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=azul
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

            cell.border = borda

        ws.row_dimensions[1].height = 25

        # ----------------------------------------------
        # DADOS
        # ----------------------------------------------

        for row in ws.iter_rows(
            min_row=2
        ):

            for cell in row:

                cell.border = borda

                cell.alignment = Alignment(
                    vertical="center"
                )

        # ----------------------------------------------
        # COLUNAS MONETÁRIAS
        # ----------------------------------------------

        colunas_monetarias = set()

        for cell in ws[1]:

            if cell.value in [
                "Valor",
                "Faturamento"
            ]:

                colunas_monetarias.add(
                    cell.column
                )

        for numero_coluna in colunas_monetarias:

            for linha in range(
                2,
                ws.max_row + 1
            ):

                ws.cell(
                    linha,
                    numero_coluna
                ).number_format = (
                    'R$ #,##0.00'
                )

        # ----------------------------------------------
        # LARGURA DAS COLUNAS
        # ----------------------------------------------

        for coluna in ws.columns:

            numero = coluna[0].column

            letra = get_column_letter(numero)

            maior = 0

            eh_monetaria = (
                numero in colunas_monetarias
            )

            for cell in coluna:

                valor = (
                    ""
                    if cell.value is None
                    else str(cell.value)
                )

                if eh_monetaria:

                    try:

                        valor = (
                            f"R$ "
                            f"{float(cell.value):,.2f}"
                        )

                    except (
                        ValueError,
                        TypeError
                    ):
                        pass

                maior = max(
                    maior,
                    len(valor)
                )

            ws.column_dimensions[
                letra
            ].width = min(
                max(maior + 4, 10),
                40
            )

        # ----------------------------------------------
        # CONGELAR CABEÇALHO
        # ----------------------------------------------

        ws.freeze_panes = "A2"

        # ----------------------------------------------
        # TABELA
        # ----------------------------------------------

        if ws.max_row >= 2:

            nome = (
                "Tabela_"
                + str(
                    abs(
                        hash(ws.title)
                    )
                )[:8]
            )

            tabela = Table(
                displayName=nome,
                ref=ws.dimensions
            )

            estilo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False
            )

            tabela.tableStyleInfo = estilo

            ws.add_table(tabela)


# ==========================================================
# CRIAR DASHBOARD
# ==========================================================

def criar_dashboard(
    wb,
    df,
    registros_originais,
    duplicados
):

    ws = wb["Resumo"]

    # ----------------------------------------------
    # LIMPA A ABA
    # ----------------------------------------------

    for row in ws.iter_rows():

     for cell in row:

        cell.value = None

    # ----------------------------------------------
    # CORES
    # ----------------------------------------------

    azul = "1F4E78"
    azul_claro = "D9EAF7"
    verde = "548235"
    verde_claro = "E2F0D9"
    branco = "FFFFFF"
    cinza = "666666"

    # ----------------------------------------------
    # TÍTULO
    # ----------------------------------------------

    ws.merge_cells(
        "A1:F1"
    )

    ws["A1"] = (
        "RELATÓRIO DE ANÁLISE DE PLANILHA"
    )

    ws["A1"].font = Font(
        bold=True,
        size=18,
        color=branco
    )

    ws["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor=azul
    )

    ws["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    ws.row_dimensions[1].height = 35

    # ----------------------------------------------
    # SUBTÍTULO
    # ----------------------------------------------

    ws.merge_cells(
        "A2:F2"
    )

    ws["A2"] = (
        "Resumo automático do processamento"
    )

    ws["A2"].font = Font(
        italic=True,
        color=cinza
    )

    ws["A2"].alignment = Alignment(
        horizontal="center"
    )

    # ----------------------------------------------
    # INDICADORES
    # ----------------------------------------------

    cards = [
        (
            "A4:B4",
            "A5:B6",
            "REGISTROS ORIGINAIS",
            registros_originais
        ),

        (
            "C4:D4",
            "C5:D6",
            "REGISTROS FINAIS",
            len(df)
        ),

        (
            "E4:F4",
            "E5:F6",
            "DUPLICADOS REMOVIDOS",
            duplicados
        )
    ]

    for (
        area_titulo,
        area_valor,
        titulo,
        valor
    ) in cards:

        ws.merge_cells(area_titulo)
        ws.merge_cells(area_valor)

        celula_titulo = ws[
            area_titulo.split(":")[0]
        ]

        celula_valor = ws[
            area_valor.split(":")[0]
        ]

        celula_titulo.value = titulo
        celula_valor.value = valor

        celula_titulo.font = Font(
            bold=True,
            color=branco
        )

        celula_titulo.fill = PatternFill(
            fill_type="solid",
            fgColor=azul
        )

        celula_titulo.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        celula_valor.font = Font(
            bold=True,
            size=20,
            color=azul
        )

        celula_valor.fill = PatternFill(
            fill_type="solid",
            fgColor=azul_claro
        )

        celula_valor.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # ----------------------------------------------
    # FATURAMENTO
    # ----------------------------------------------

    faturamento = calcular_faturamento(df)

    ws.merge_cells(
        "A9:F9"
    )

    ws["A9"] = "FATURAMENTO TOTAL"

    ws["A9"].font = Font(
        bold=True,
        color=branco
    )

    ws["A9"].fill = PatternFill(
        fill_type="solid",
        fgColor=verde
    )

    ws["A9"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    ws.merge_cells(
        "A10:F11"
    )

    ws["A10"] = faturamento

    ws["A10"].font = Font(
        bold=True,
        size=24,
        color=verde
    )

    ws["A10"].fill = PatternFill(
        fill_type="solid",
        fgColor=verde_claro
    )

    ws["A10"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    ws["A10"].number_format = (
        'R$ #,##0.00'
    )

    # ----------------------------------------------
    # MENSAGEM
    # ----------------------------------------------

    ws.merge_cells(
        "A13:F13"
    )

    ws["A13"] = (
        "Relatório gerado automaticamente."
    )

    ws["A13"].font = Font(
        italic=True,
        color=cinza
    )

    ws["A13"].alignment = Alignment(
        horizontal="center"
    )

    # ----------------------------------------------
    # LARGURA
    # ----------------------------------------------

    for coluna in [
        "A",
        "B",
        "C",
        "D",
        "E",
        "F"
    ]:

        ws.column_dimensions[
            coluna
        ].width = 18

    # ----------------------------------------------
    # ALTURA
    # ----------------------------------------------

    ws.row_dimensions[5].height = 30
    ws.row_dimensions[6].height = 30
    ws.row_dimensions[10].height = 30
    ws.row_dimensions[11].height = 30

    # ----------------------------------------------
    # SEM CONGELAMENTO
    # ----------------------------------------------

    ws.freeze_panes = None


# ==========================================================
# FORMATAÇÃO FINAL
# ==========================================================

def formatar_relatorio(
    caminho_saida,
    df,
    registros_originais,
    duplicados
):

    wb = load_workbook(
        caminho_saida
    )

    # Formata as abas de dados
    formatar_abas(wb)

    # Cria dashboard
    criar_dashboard(
        wb,
        df,
        registros_originais,
        duplicados
    )

    # ----------------------------------------------
    # RESUMO COMO PRIMEIRA ABA
    # ----------------------------------------------

    abas = wb.sheetnames

    if "Resumo" in abas:

        abas.remove("Resumo")

        abas.insert(
            0,
            "Resumo"
        )

        wb._sheets = [
            wb[nome]
            for nome in abas
        ]

    # Salva
    wb.save(
        caminho_saida
    )


# ==========================================================
# PROGRAMA PRINCIPAL
# ==========================================================

def main():

    print("=" * 55)
    print(
        "        AUTOMATIZADOR DE PLANILHAS"
    )
    print("=" * 55)

    caminho = input(
        "\nDigite o caminho do arquivo: "
    ).strip()

    try:

        # ----------------------------------------------
        # 1
        # ----------------------------------------------

        print(
            "\n[1/4] Carregando arquivo..."
        )

        df = carregar_arquivo(
            caminho
        )

        # ----------------------------------------------
        # 2
        # ----------------------------------------------

        print(
            "[2/4] Limpando dados..."
        )

        (
            df_limpo,
            registros_originais,
            duplicados
        ) = limpar_dados(df)

        # ----------------------------------------------
        # 3
        # ----------------------------------------------

        print(
            "[3/4] Gerando relatório..."
        )

        caminho_saida = criar_relatorio(
            df_limpo,
            registros_originais,
            duplicados,
            "saida/relatorio.xlsx"
        )

        formatar_relatorio(
            caminho_saida,
            df_limpo,
            registros_originais,
            duplicados
        )

        # ----------------------------------------------
        # 4
        # ----------------------------------------------

        print(
            "[4/4] Finalizando..."
        )

        print("\n" + "=" * 55)
        print(
            "          PROCESSAMENTO CONCLUÍDO"
        )
        print("=" * 55)

        print(
            f"\nRegistros originais: "
            f"{registros_originais}"
        )

        print(
            f"Registros finais: "
            f"{len(df_limpo)}"
        )

        print(
            f"Duplicados removidos: "
            f"{duplicados}"
        )

        print(
            f"Colunas: "
            f"{len(df_limpo.columns)}"
        )

        print(
            "\nRelatório criado em:"
        )

        print(
            caminho_saida.resolve()
        )

    except Exception as erro:

        print("\n" + "=" * 55)
        print(
            "ERRO DURANTE O PROCESSAMENTO"
        )
        print("=" * 55)

        print(
            f"\n{erro}"
        )


# ==========================================================
# INICIAR
# ==========================================================

if __name__ == "__main__":
    main()